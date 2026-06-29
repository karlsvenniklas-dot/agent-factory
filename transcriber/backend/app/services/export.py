"""Export transcript in multiple formats."""
from __future__ import annotations

from io import BytesIO
from typing import TYPE_CHECKING

if TYPE_CHECKING:
    from ..models import Meeting, Segment, Speaker


def _format_ts_srt(seconds: float) -> str:
    h = int(seconds // 3600)
    m = int((seconds % 3600) // 60)
    s = int(seconds % 60)
    ms = int((seconds - int(seconds)) * 1000)
    return f"{h:02d}:{m:02d}:{s:02d},{ms:03d}"


def _format_ts_vtt(seconds: float) -> str:
    return _format_ts_srt(seconds).replace(",", ".")


def export_srt(meeting: "Meeting", segments: list["Segment"], speakers: dict[str, "Speaker"]) -> bytes:
    lines = []
    for i, seg in enumerate(segments, 1):
        spk = speakers.get(seg.speaker_id)
        name = spk.name if spk else "Okänd"
        lines.append(str(i))
        lines.append(f"{_format_ts_srt(seg.start_time)} --> {_format_ts_srt(seg.end_time)}")
        lines.append(f"<{name}> {seg.text}")
        lines.append("")
    return "\n".join(lines).encode("utf-8")


def export_vtt(meeting: "Meeting", segments: list["Segment"], speakers: dict[str, "Speaker"]) -> bytes:
    lines = ["WEBVTT", ""]
    for i, seg in enumerate(segments, 1):
        spk = speakers.get(seg.speaker_id)
        name = spk.name if spk else "Okänd"
        lines.append(f"{i}")
        lines.append(f"{_format_ts_vtt(seg.start_time)} --> {_format_ts_vtt(seg.end_time)}")
        lines.append(f"<{name}> {seg.text}")
        lines.append("")
    return "\n".join(lines).encode("utf-8")


def export_text(meeting: "Meeting", segments: list["Segment"], speakers: dict[str, "Speaker"]) -> bytes:
    lines = [f"# {meeting.title}", ""]
    current_speaker = None
    for seg in segments:
        spk = speakers.get(seg.speaker_id)
        name = spk.name if spk else "Okänd"
        if name != current_speaker:
            lines.append(f"\n{name}:")
            current_speaker = name
        lines.append(f"  {seg.text}")
    return "\n".join(lines).encode("utf-8")


def export_markdown(meeting: "Meeting", segments: list["Segment"], speakers: dict[str, "Speaker"]) -> bytes:
    lines = [f"# {meeting.title}", ""]
    current_speaker = None
    for seg in segments:
        spk = speakers.get(seg.speaker_id)
        name = spk.name if spk else "Okänd"
        ts = f"[{_format_ts_vtt(seg.start_time)}]"
        if name != current_speaker:
            lines.append(f"\n**{name}** {ts}")
            current_speaker = name
        lines.append(f"> {seg.text}")
    return "\n".join(lines).encode("utf-8")


def export_json(meeting: "Meeting", segments: list["Segment"], speakers: dict[str, "Speaker"]) -> bytes:
    import json
    data = {
        "id": meeting.id,
        "title": meeting.title,
        "duration": meeting.duration_seconds,
        "language": meeting.language,
        "speakers": [
            {"id": s.id, "label": s.label, "name": s.name, "color": s.color}
            for s in speakers.values()
        ],
        "segments": [
            {
                "id": seg.id,
                "speaker_id": seg.speaker_id,
                "speaker_name": speakers[seg.speaker_id].name if seg.speaker_id and seg.speaker_id in speakers else None,
                "start": seg.start_time,
                "end": seg.end_time,
                "text": seg.text,
                "confidence": seg.confidence,
            }
            for seg in segments
        ],
    }
    return json.dumps(data, ensure_ascii=False, indent=2).encode("utf-8")


def export_docx(meeting: "Meeting", segments: list["Segment"], speakers: dict[str, "Speaker"]) -> bytes:
    from docx import Document  # type: ignore
    from docx.shared import RGBColor, Pt  # type: ignore

    doc = Document()
    doc.add_heading(meeting.title, 0)

    current_speaker = None
    current_para = None
    for seg in segments:
        spk = speakers.get(seg.speaker_id)
        name = spk.name if spk else "Okänd"
        ts = f"[{_format_ts_vtt(seg.start_time)}]"
        if name != current_speaker:
            p = doc.add_paragraph()
            run = p.add_run(f"{name} {ts}\n")
            run.bold = True
            current_speaker = name
            current_para = doc.add_paragraph()
        else:
            current_para = doc.add_paragraph()
        current_para.add_run(seg.text)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def export_pdf(meeting: "Meeting", segments: list["Segment"], speakers: dict[str, "Speaker"]) -> bytes:
    from reportlab.lib.pagesizes import A4  # type: ignore
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle  # type: ignore
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer  # type: ignore
    from reportlab.lib.units import cm  # type: ignore

    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=2*cm, rightMargin=2*cm, topMargin=2*cm, bottomMargin=2*cm)
    styles = getSampleStyleSheet()
    bold = ParagraphStyle("bold", parent=styles["Normal"], fontName="Helvetica-Bold", fontSize=10)
    normal = ParagraphStyle("normal", parent=styles["Normal"], fontSize=10, leading=14)

    story = [Paragraph(meeting.title, styles["Title"]), Spacer(1, 0.5*cm)]
    current_speaker = None
    for seg in segments:
        spk = speakers.get(seg.speaker_id)
        name = spk.name if spk else "Okänd"
        ts = _format_ts_vtt(seg.start_time)
        if name != current_speaker:
            story.append(Spacer(1, 0.3*cm))
            story.append(Paragraph(f"{name} [{ts}]", bold))
            current_speaker = name
        story.append(Paragraph(seg.text, normal))

    doc.build(story)
    return buf.getvalue()


EXPORTERS = {
    "srt": (export_srt, "text/srt", ".srt"),
    "vtt": (export_vtt, "text/vtt", ".vtt"),
    "txt": (export_text, "text/plain", ".txt"),
    "md": (export_markdown, "text/markdown", ".md"),
    "json": (export_json, "application/json", ".json"),
    "docx": (export_docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document", ".docx"),
    "pdf": (export_pdf, "application/pdf", ".pdf"),
}
